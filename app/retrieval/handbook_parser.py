"""Extract a policy handbook (.docx, .pdf, .txt/.md) into citable sections.

Each section keeps the document's own heading and number exactly as written,
so every citation the system produces can be traced to real text. Tables are
split into one chunk per row (e.g. one row of the escalation table per role),
labelled with the heading the table sits under; no new numbering is invented.
"""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

# Sections that are shown in the Policy Library but not used as case evidence.
# - Contents / Document Control / Version History: document metadata.
# - Appendix A: the AI case-handling rules. They are built into every agent's
#   instructions instead of being "retrieved" as evidence.
# - Appendix B: the handbook's own synthetic example cases. Excluded so that the
#   evaluation does not simply retrieve the expected answers.
# - Official and Government Reference Sources: used by the Compliance Reviewer
#   as "where to verify" pointers, not as policy evidence.
NON_EVIDENCE_HEADINGS = {
    "cover",
    "contents",
    "document control",
    "version history",
    "important notice: academic use only",
    "official and government reference sources",
}
NON_EVIDENCE_PREFIXES = ("appendix a", "appendix b")


@dataclass
class PolicySection:
    chunk_id: str
    document: str
    section_number: str  # exactly as printed, e.g. "3.3", "14", "Appendix C"
    heading: str  # heading text without the number
    parent: str  # parent heading as printed (e.g. "3. Working Hours, ...")
    ref: str  # human-readable citation label
    text: str
    kind: str = "policy"  # policy | legal_map | escalation_table | rules | example | meta
    retrievable: bool = True
    order: int = 0
    row_label: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class ParsedDocument:
    name: str
    title: str
    sections: list[PolicySection] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


_H1_NUM = re.compile(r"^(\d+)\.\s+(.*)$")
_H2_NUM = re.compile(r"^(\d+\.\d+)\s+(.*)$")
_APPENDIX = re.compile(r"^(Appendix [A-Z])\.?\s*(.*)$")
_CASE = re.compile(r"^(Case [A-Z]\d+)\s*[—\-–]\s*(.*)$")


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60]


def _split_number(heading: str, level: int) -> tuple[str, str]:
    """Return (number, title) using only numbering already in the heading."""
    heading = heading.strip()
    if level == 2:
        m = _H2_NUM.match(heading)
        if m:
            return m.group(1), m.group(2).strip()
        m = _CASE.match(heading)
        if m:
            return m.group(1), m.group(2).strip()
    m = _H1_NUM.match(heading)
    if m:
        return m.group(1), m.group(2).strip()
    m = _APPENDIX.match(heading)
    if m:
        return m.group(1), m.group(2).strip()
    return "", heading


def _iter_block_items(doc):
    """Yield paragraphs and tables in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


class _Builder:
    def __init__(self, doc_name: str):
        self.doc_name = doc_name
        self.sections: list[PolicySection] = []
        self.h1 = "Cover"
        self.h1_number, self.h1_title = "", "Cover"
        self.h2: str | None = None
        self.h2_number, self.h2_title = "", ""
        self.buffer: list[str] = []
        self.list_counter = 0

    # -- helpers -----------------------------------------------------------
    def _current_ids(self):
        if self.h2:
            number = self.h2_number or self.h1_number
            heading = self.h2_title if self.h2_number else self.h2
            parent = self.h1
        else:
            number = self.h1_number
            heading = self.h1_title
            parent = ""
        return number, heading, parent

    def _ref(self, number: str, heading: str, parent: str, row: str = "") -> str:
        if number.startswith("Case "):
            base = f"Appendix B › {number} — {heading}"
        elif number and re.match(r"^\d", number):
            base = f"§{number} {heading}"
        elif number:
            base = f"{number} — {heading}" if heading else number
        elif parent:
            pnum, ptitle = _split_number(parent, 1)
            base = f"{pnum or ptitle} › {heading}"
        else:
            base = heading
        return f"{base} (row: {row})" if row else base

    def _classify(self, number: str, heading: str, parent: str) -> tuple[str, bool]:
        top = (parent or heading).lower()
        top_full = (parent or f"{number} {heading}").lower()
        if heading.lower() in NON_EVIDENCE_HEADINGS or top in NON_EVIDENCE_HEADINGS:
            return "meta", False
        if top_full.startswith("appendix a") or number == "Appendix A":
            return "rules", False
        if top_full.startswith("appendix b") or number.startswith("Case ") or number == "Appendix B":
            return "example", False
        if top_full.startswith("appendix c") or number == "Appendix C":
            return "legal_map", True
        return "policy", True

    def flush(self):
        text = "\n".join(t for t in self.buffer if t.strip()).strip()
        self.buffer = []
        self.list_counter = 0
        if not text:
            return
        number, heading, parent = self._current_ids()
        kind, retrievable = self._classify(number, heading, parent)
        ref = self._ref(number, heading, parent)
        chunk_id = _slug(f"{number or 'x'}-{heading}") or f"chunk-{len(self.sections)}"
        if any(s.chunk_id == chunk_id for s in self.sections):
            chunk_id = f"{chunk_id}-{len(self.sections)}"
        self.sections.append(
            PolicySection(
                chunk_id=chunk_id,
                document=self.doc_name,
                section_number=number,
                heading=heading,
                parent=parent,
                ref=ref,
                text=text,
                kind=kind,
                retrievable=retrievable,
                order=len(self.sections),
            )
        )

    def add_table(self, rows: list[list[str]]):
        """Add one chunk per table row so each row is individually citable."""
        self.flush()
        if not rows:
            return
        number, heading, parent = self._current_ids()
        kind, retrievable = self._classify(number, heading, parent)
        header = rows[0]
        is_header = len(rows) > 1 and len(header) > 1
        if heading.lower() == "document control" or not is_header:
            # Key/value table: keep as one chunk.
            text = "\n".join(" | ".join(c for c in r if c) for r in rows)
            self.buffer = [text]
            self.flush()
            return
        if "role" in header[0].lower():
            kind = "escalation_table"
        for r in rows[1:]:
            label = r[0].strip()
            parts = [f"{h}: {c}" for h, c in zip(header, r) if c.strip()]
            text = "\n".join(parts)
            base_slug = _slug(f"{number or 'x'}-{heading}")[:36]
            chunk_id = f"{base_slug}-row-{_slug(label)[:28]}"
            if any(s.chunk_id == chunk_id for s in self.sections):
                chunk_id = f"{chunk_id}-{len(self.sections)}"
            self.sections.append(
                PolicySection(
                    chunk_id=chunk_id,
                    document=self.doc_name,
                    section_number=number,
                    heading=heading,
                    parent=parent,
                    ref=self._ref(number, heading, parent, row=label),
                    text=text,
                    kind=kind,
                    retrievable=retrievable,
                    order=len(self.sections),
                    row_label=label,
                )
            )

    def heading1(self, text: str):
        self.flush()
        self.h1 = text.strip()
        self.h1_number, self.h1_title = _split_number(text, 1)
        self.h2 = None
        self.h2_number = self.h2_title = ""

    def heading2(self, text: str):
        self.flush()
        self.h2 = text.strip()
        self.h2_number, self.h2_title = _split_number(text, 2)

    def paragraph(self, text: str, style: str):
        text = text.strip()
        if not text:
            return
        if style.startswith("List Number"):
            self.list_counter += 1
            text = f"{self.list_counter}. {text}"
        elif style.startswith("List"):
            text = f"• {text}"
        self.buffer.append(text)


def parse_docx(path: Path) -> ParsedDocument:
    doc = Document(str(path))
    builder = _Builder(path.name)
    title_parts: list[str] = []
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            style = (block.style.name if block.style is not None else "Normal") or "Normal"
            text = block.text
            if style == "Title":
                title_parts.append(text.replace("\n", " ").strip())
                builder.paragraph(text.replace("\n", " "), "Normal")
            elif style.startswith("Heading 1"):
                builder.heading1(text)
            elif style.startswith("Heading 2") or style.startswith("Heading 3"):
                builder.heading2(text)
            else:
                builder.paragraph(text, style)
        else:
            rows = []
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                # merged cells repeat; de-duplicate consecutive repeats
                dedup = []
                for c in cells:
                    if not dedup or dedup[-1] != c:
                        dedup.append(c)
                rows.append(dedup)
            builder.add_table(rows)
    builder.flush()
    title = " ".join(t for t in title_parts if t) or path.stem
    parsed = ParsedDocument(name=path.name, title=title, sections=builder.sections)
    if not any(s.retrievable for s in parsed.sections):
        parsed.warnings.append("No headed policy sections were found; check that the document uses Heading 1/Heading 2 styles.")
    return parsed


_TEXT_HEADING = re.compile(r"^(#{1,2})\s+(.*)$|^(\d+(?:\.\d+)?)\.?\s+([A-Z].{2,80})$")


def parse_text(path: Path, text: str | None = None) -> ParsedDocument:
    """Parse plain text / markdown / extracted PDF text using numbered headings."""
    if text is None:
        text = path.read_text(encoding="utf-8", errors="replace")
    builder = _Builder(path.name)
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        m = _TEXT_HEADING.match(s)
        if m and m.group(1):
            (builder.heading1 if len(m.group(1)) == 1 else builder.heading2)(m.group(2))
        elif m and m.group(3) and len(s) < 90 and not s.endswith("."):
            (builder.heading2 if "." in m.group(3) else builder.heading1)(s)
        else:
            builder.paragraph(s, "Normal")
    builder.flush()
    parsed = ParsedDocument(name=path.name, title=path.stem, sections=builder.sections)
    if len(parsed.sections) <= 1:
        parsed.warnings.append("Few headings detected; citations will be coarse.")
    return parsed


def parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return parse_text(path, text)


def parse_policy_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix in {".txt", ".md"}:
        return parse_text(path)
    raise ValueError(f"Unsupported policy file type: {suffix}")
