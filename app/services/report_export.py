"""Export the draft HR case report as DOCX (python-docx) and PDF (reportlab).

Both formats are built from the same section list, so they always match the
on-screen report.
"""
from __future__ import annotations

import io
from xml.sax.saxutils import escape

DRAFT_BANNER = "DRAFT FOR HR REVIEW — FICTIONAL ACADEMIC PROTOTYPE — NOT AN EMPLOYMENT DECISION"


def _fmt_list(items):
    return [str(x) for x in items if str(x).strip()]


def report_sections(case: dict) -> list[tuple[str, list[tuple]]]:
    r = case["report"]
    gate = case.get("gate") or {}
    review = r.get("human_review", {})
    out: list[tuple[str, list[tuple]]] = []

    out.append(("1. Case ID and category", [("table", ["Field", "Value"], [
        ["Case ID", r["case_id"]],
        ["Title", r.get("title") or ""],
        ["Category", r.get("category_label") or ""],
        ["Urgency", r.get("urgency") or ""],
        ["Route", r.get("route_label") or ""],
        ["Access", "RESTRICTED — designated HR/IC only" if r.get("restricted_access") else "Standard HR access"],
        ["Analysis mode", "Live AI" if case.get("mode") == "live" else "Demo mode (no language model; narrative simulated)"],
        ["Generated", r.get("generated_at", "")],
    ])] + ([("note", r["category_note"])] if r.get("category_note") else [])))

    out.append(("2. Executive summary", [("p", r.get("executive_summary", ""))]))

    rc = r["reported_concern"]
    blocks = [("note", rc.get("label", "Reported account — not verified")), ("p", rc.get("summary", ""))]
    if rc.get("statements"):
        blocks.append(("bullets", [f"The reporter states: “{s}”" for s in rc["statements"]]))
    if rc.get("requested_outcome"):
        blocks.append(("p", f"Requested outcome (as stated): {rc['requested_outcome']}"))
    out.append(("3. Employee's reported concern", blocks))

    cf = r["confirmed_facts"]
    blocks = [("note", cf.get("note", "")), ("bullets", _fmt_list(cf.get("case_record_facts", [])))]
    for d in cf.get("documentary_evidence", []):
        blocks.append(("p", f"Document: {d['name']} — {d['status']}."))
        if d.get("notable_lines"):
            blocks.append(("bullets", [f"Line in document: “{l}”" for l in d["notable_lines"]]))
    out.append(("4. Confirmed facts and available evidence", blocks))

    blocks = []
    if r.get("policy_provisions"):
        rows = [[p["ref"], p.get("quote", ""), p.get("how_it_applies") or "", p.get("source_type", "")] for p in r["policy_provisions"]]
        blocks.append(("table", ["Section", "Handbook text (verbatim)", "Relevance", "Source type"], rows))
        doc = r["policy_provisions"][0].get("document")
        blocks.append(("note", f"Source document: {doc}"))
    else:
        blocks.append(("p", "No supporting handbook passage was retrieved for this concern. Human review must decide which policy, if any, applies."))
    for g in r.get("policy_gaps", []):
        blocks.append(("note", g))
    for n in r.get("policy_vs_law_notes", []):
        blocks.append(("note", n))
    out.append(("5. Relevant HR policy provisions", blocks))

    mi = r.get("missing_information", [])
    out.append(("6. Missing information", [("table", ["Item", "Priority", "Basis"], [[m["item"], m.get("severity", ""), m.get("basis", "")] for m in mi])] if mi else [("p", "No missing items identified.")]))

    comp = r["compliance"]
    blocks = [("note", comp.get("disclaimer", ""))]
    if comp.get("statutory_route"):
        blocks.append(("p", comp["statutory_route"]))
    rows = [[i["topic"], i.get("framework_to_verify") or "Not mapped in handbook — Legal to identify", i.get("what_hr_must_confirm") or "", i["status"]] for i in comp.get("issues", [])]
    if rows:
        blocks.append(("table", ["Topic", "Framework to verify (handbook Appendix C)", "What HR must confirm", "Status"], rows))
    facts = comp.get("applicability_facts", [])
    if facts:
        blocks.append(("table", ["Applicability fact", "Value on file", "Why needed"], [[f["fact"], f["value"] or "MISSING", f["why_needed"]] for f in facts]))
    if comp.get("verification_sources"):
        blocks.append(("bullets", [f"{s['label']}: {s['url']}" for s in comp["verification_sources"]]))
    out.append(("7. Compliance issues requiring verification", blocks))

    out.append(("8. Suggested follow-up questions", [("bullets", _fmt_list(r.get("follow_up_questions", [])))]))

    steps = r.get("next_steps", [])
    blocks = [("table", ["Step", "Responsible", "Basis"], [[s["step"], s.get("owner", ""), s.get("ref", "")] for s in steps])]
    if r.get("ai_suggested_steps"):
        blocks.append(("note", "Additional AI-suggested steps (for reviewer consideration):"))
        blocks.append(("bullets", r["ai_suggested_steps"]))
    if r.get("conflicts_of_interest"):
        blocks.append(("note", "Conflicts of interest:"))
        blocks.append(("bullets", r["conflicts_of_interest"]))
    if r.get("evidence_checklist"):
        blocks.append(("table", ["Evidence checklist item", "Status", "Basis"], [[e["item"], e["status"], e["basis"]] for e in r["evidence_checklist"]]))
    if r.get("responsible_functions"):
        blocks.append(("p", "Responsible HR functions: " + ", ".join(r["responsible_functions"])))
    out.append(("9. Recommended procedural next steps", blocks))

    esc = r["escalation"]
    blocks = []
    if esc.get("urgent_actions"):
        blocks.append(("note", "URGENT — act now, do not wait for this report:"))
        blocks.append(("bullets", esc["urgent_actions"]))
    blocks.append(("table", ["Escalate to", "Priority", "Reason", "Basis"], [[e["to"], e["priority"], e["reason"], "; ".join(b["ref"] for b in e.get("basis", []))] for e in esc.get("items", [])]))
    out.append(("10. Escalation requirements", blocks))

    blocks = [("bullets", _fmt_list(r.get("limitations", [])))]
    if r.get("unresolved_questions"):
        blocks.append(("note", "Unresolved questions:"))
        blocks.append(("bullets", r["unresolved_questions"]))
    if gate:
        blocks.append(("p", f"Quality & Risk Gate: {gate.get('verdict_label', '')}"))
        blocks.append(("table", ["Check", "Result", "Detail"], [[c["check"], c["status"].upper(), c["detail"]] for c in gate.get("checks", [])]))
        if gate.get("corrections"):
            blocks.append(("bullets", gate["corrections"]))
    out.append(("11. Limitations and uncertainty", blocks))

    status = {"pending": "PENDING — not approved", "approved": "APPROVED (draft accepted for further HR action)", "returned": "RETURNED for rework"}.get(review.get("status"), review.get("status"))
    rows = [["Status", status], ["Reviewer", review.get("reviewer") or "—"], ["Decided", review.get("decided_at") or "—"], ["Notes", review.get("notes") or "—"]]
    out.append(("12. Human review and approval status", [("table", ["Field", "Value"], rows), ("note", "Approval confirms the draft for further HR handling only. It is not a finding on the allegation and not a disciplinary decision.")]))
    return out


# ---------------------------------------------------------------------------
def build_docx(case: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    navy = RGBColor(0x14, 0x2B, 0x4F)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    banner = doc.add_paragraph()
    run = banner.add_run(DRAFT_BANNER)
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xA1, 0x2C, 0x2C)
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h = doc.add_heading(f"PeopleGuard AI — HR Case Report {case['id']}", level=0)
    for r in h.runs:
        r.font.color.rgb = navy
    doc.add_paragraph("Northstar Retail Services Pvt. Ltd. (fictional). Source handbook: PeopleGuard_Fictional_HR_Policy_Book_India.docx")

    for title, blocks in report_sections(case):
        hh = doc.add_heading(title, level=1)
        for r in hh.runs:
            r.font.color.rgb = navy
        for b in blocks:
            kind = b[0]
            if kind == "p" and b[1]:
                doc.add_paragraph(b[1])
            elif kind == "note" and b[1]:
                p = doc.add_paragraph()
                rr = p.add_run(b[1])
                rr.italic = True
                rr.font.size = Pt(9)
            elif kind == "bullets":
                for item in b[1]:
                    doc.add_paragraph(item, style="List Bullet")
            elif kind == "table":
                headers, rows = b[1], b[2]
                if not rows:
                    continue
                t = doc.add_table(rows=1, cols=len(headers))
                t.style = "Light Grid Accent 1"
                for i, hdr in enumerate(headers):
                    t.rows[0].cells[i].text = hdr
                for row in rows:
                    cells = t.add_row().cells
                    for i, val in enumerate(row):
                        cells[i].text = str(val)
                doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(case: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    navy = colors.HexColor("#142B4F")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10)
    note = ParagraphStyle("note", parent=body, fontName="Helvetica-Oblique", textColor=colors.HexColor("#4A5568"))
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=17, textColor=navy, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=navy, spaceBefore=10, spaceAfter=4)
    banner = ParagraphStyle("banner", parent=small, alignment=TA_CENTER, textColor=colors.HexColor("#A12C2C"), fontName="Helvetica-Bold")

    def P(text, st=body):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), st)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"PeopleGuard AI case report {case['id']}", author="PeopleGuard AI (academic prototype)")
    width = A4[0] - 32 * mm
    story = [P(DRAFT_BANNER, banner), Spacer(1, 4), P(f"PeopleGuard AI — HR Case Report {case['id']}", h1),
             P("Northstar Retail Services Pvt. Ltd. (fictional). Source handbook: PeopleGuard_Fictional_HR_Policy_Book_India.docx", small), Spacer(1, 6)]

    for title, blocks in report_sections(case):
        story.append(P(title, h2))
        for b in blocks:
            kind = b[0]
            if kind == "p" and b[1]:
                story.append(P(b[1]))
                story.append(Spacer(1, 3))
            elif kind == "note" and b[1]:
                story.append(P(b[1], note))
                story.append(Spacer(1, 3))
            elif kind == "bullets" and b[1]:
                story.append(ListFlowable([ListItem(P(i), leftIndent=10) for i in b[1]], bulletType="bullet", start="•", leftIndent=10))
                story.append(Spacer(1, 3))
            elif kind == "table" and b[2]:
                headers, rows = b[1], b[2]
                n = len(headers)
                col = [width / n] * n
                if n == 4:
                    col = [width * 0.18, width * 0.40, width * 0.27, width * 0.15]
                elif n == 3:
                    col = [width * 0.5, width * 0.2, width * 0.3]
                elif n == 2:
                    col = [width * 0.25, width * 0.75]
                data = [[P(hd, small) for hd in headers]] + [[P(c, small) for c in row] for row in rows]
                t = Table(data, colWidths=col, repeatRows=1)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF5")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D2E0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(t)
                story.append(Spacer(1, 5))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#718096"))
        canvas.drawString(16 * mm, 8 * mm, f"{case['id']} — draft for HR review — fictional data")
        canvas.drawRightString(A4[0] - 16 * mm, 8 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
