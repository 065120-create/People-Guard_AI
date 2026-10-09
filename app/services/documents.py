"""Text extraction for documents attached to a case (PDF, DOCX, TXT)."""
from __future__ import annotations

import io

from app.services.rules import find_pii

ALLOWED = {".pdf", ".docx", ".txt", ".md"}
MAX_CHARS = 20000


class DocumentError(ValueError):
    pass


def extract_text(filename: str, data: bytes, max_bytes: int) -> dict:
    name = (filename or "upload").replace("\\", "/").split("/")[-1][:120]
    suffix = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if suffix not in ALLOWED:
        raise DocumentError(f"Unsupported file type '{suffix or 'none'}'. Upload a PDF, DOCX or TXT file.")
    if len(data) > max_bytes:
        raise DocumentError(f"File is too large ({len(data) // 1024} KB). The limit is {max_bytes // 1024} KB.")
    if not data:
        raise DocumentError("The uploaded file is empty.")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif suffix == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except DocumentError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Could not read the file ({type(exc).__name__}). Is it a valid {suffix} file?") from exc
    text = text.replace("\x00", "").strip()
    truncated = len(text) > MAX_CHARS
    text = text[:MAX_CHARS]
    if not text:
        raise DocumentError("No readable text was found in the file (scanned PDFs are not supported).")
    return {"name": name, "text": text, "chars": len(text), "truncated": truncated, "pii": find_pii(text)}
